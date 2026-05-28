#!/usr/bin/env python3
"""
Embed MP4 videos into the PowerPoint presentation
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import os

PPTX_PATH = '/Users/moderndavinci/Desktop/lightning-pass-beach-club-pitch.pptx'
VIDEOS_DIR = '/Users/moderndavinci/Desktop/Applications/Lightning-Project/Lightning-App/screenshots/videos'

# Video to slide mapping (0-indexed slide numbers)
# Slides 16-20 are the demo slides (indices 15-19)
VIDEO_MAPPING = [
    {
        'video': 'customer-flow.mp4',
        'slide_index': 15,  # Slide 16: Customer RSVP Flow
        'position': (Inches(9.5), Inches(1.3)),
        'size': (Inches(3), Inches(5.5)),
    },
    {
        'video': 'dj-experience.mp4',
        'slide_index': 16,  # Slide 17: DJ Song Request System
        'position': (Inches(9.5), Inches(1.3)),
        'size': (Inches(3.5), Inches(2)),
    },
    {
        'video': 'venue-dashboard.mp4',
        'slide_index': 17,  # Slide 18: Venue Admin Dashboard
        'position': (Inches(9.5), Inches(5.5)),
        'size': (Inches(3.5), Inches(2)),
    },
    {
        'video': 'white-label-demo.mp4',
        'slide_index': 18,  # Slide 19: White-Label Comparison
        'position': (Inches(10), Inches(4)),
        'size': (Inches(2.5), Inches(2.5)),
    },
    {
        'video': 'bouncer-scan.mp4',
        'slide_index': 19,  # Slide 20: Door Scanner
        'position': (Inches(9.5), Inches(4)),
        'size': (Inches(3.5), Inches(2)),
    },
]

def main():
    print("\n🎬 Embedding videos into PowerPoint...\n")

    prs = Presentation(PPTX_PATH)
    print(f"   Loaded presentation with {len(prs.slides)} slides")

    embedded = 0
    for mapping in VIDEO_MAPPING:
        video_path = os.path.join(VIDEOS_DIR, mapping['video'])
        slide_idx = mapping['slide_index']

        if not os.path.exists(video_path):
            print(f"   ⚠️  {mapping['video']} not found, skipping")
            continue

        if slide_idx >= len(prs.slides):
            print(f"   ⚠️  Slide {slide_idx + 1} doesn't exist, skipping")
            continue

        slide = prs.slides[slide_idx]
        left, top = mapping['position']
        width, height = mapping['size']

        # Add video to slide
        try:
            slide.shapes.add_movie(
                video_path,
                left, top,
                width, height,
                mime_type='video/mp4'
            )
            print(f"   ✅ Embedded {mapping['video']} on slide {slide_idx + 1}")
            embedded += 1
        except Exception as e:
            print(f"   ❌ Failed to embed {mapping['video']}: {e}")

    # Save the presentation
    prs.save(PPTX_PATH)
    print(f"\n💾 Saved presentation with {embedded} embedded videos")
    print(f"   File: {PPTX_PATH}\n")

if __name__ == "__main__":
    main()
