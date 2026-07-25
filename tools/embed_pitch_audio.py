#!/usr/bin/env python3
"""Generate ElevenLabs narration for the MS5 pitch script and embed it in the deck.

Reads the narration text from the pitch-script markdown (everything between the
[START RECORDING] and [END RECORDING] markers), synthesizes it with the
ElevenLabs text-to-speech API, and embeds the resulting MP3 on a slide of the
PowerPoint as a click-to-play audio object.

Usage:
    export ELEVENLABS_API_KEY=...
    python3 tools/embed_pitch_audio.py \
        --script PSY374_MS5_Pitch_Script.md \
        --deck PSY374_MS5_EmploymentPlan_Lerner.pptx \
        --voice-id I355zzIYuOq3ZXZ6D7wC \
        --slide 14 \
        --out PSY374_MS5_EmploymentPlan_Lerner_WITH_AUDIO.pptx

The API key is taken from the ELEVENLABS_API_KEY environment variable or the
--api-key flag. Never commit the key itself to the repository.
"""

import argparse
import json
import os
import re
import sys
import urllib.request

from pptx import Presentation
from pptx.util import Inches, Pt

API_URL = "https://api.elevenlabs.io/v1/text-to-speech/{voice_id}?output_format=mp3_44100_128"


def extract_narration(md_path):
    text = open(md_path, encoding="utf-8").read()
    m = re.search(r"\[START RECORDING\]\*{0,2}(.*?)\*{0,2}\[END RECORDING",
                  text, re.DOTALL)
    if not m:
        sys.exit("Could not find [START RECORDING] ... [END RECORDING] markers")
    body = m.group(1)
    # Drop markdown emphasis markers and collapse blank lines
    body = body.replace("**", "").strip()
    paragraphs = [p.strip() for p in body.split("\n\n") if p.strip()]
    return "\n\n".join(paragraphs)


def synthesize(text, voice_id, api_key, mp3_path):
    payload = {
        "text": text,
        "model_id": "eleven_multilingual_v2",
        "voice_settings": {"stability": 0.5, "similarity_boost": 0.75},
    }
    req = urllib.request.Request(
        API_URL.format(voice_id=voice_id),
        data=json.dumps(payload).encode(),
        headers={"xi-api-key": api_key, "Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            audio = resp.read()
    except urllib.error.HTTPError as e:
        sys.exit(f"ElevenLabs API error {e.code}: {e.read().decode(errors='replace')}")
    with open(mp3_path, "wb") as f:
        f.write(audio)
    return mp3_path


def find_slide_audio_target(deck_path, slide_number):
    """Return the package path of the audio media the slide references, if any."""
    import xml.dom.minidom
    import zipfile
    from posixpath import normpath

    rels_name = f"ppt/slides/_rels/slide{slide_number}.xml.rels"
    with zipfile.ZipFile(deck_path) as z:
        if rels_name not in z.namelist():
            return None
        dom = xml.dom.minidom.parseString(z.read(rels_name))
    for rel in dom.getElementsByTagName("Relationship"):
        if rel.getAttribute("Type").endswith("/audio"):
            target = rel.getAttribute("Target")
            return normpath("ppt/slides/" + target)
    return None


def replace_media_part(deck_path, part_name, mp3_path, out_path):
    """Rewrite the package with the given media part's bytes replaced."""
    import shutil
    import zipfile

    new_bytes = open(mp3_path, "rb").read()
    tmp = out_path + ".tmp"
    with zipfile.ZipFile(deck_path) as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = new_bytes if item.filename == part_name else zin.read(item)
            zout.writestr(item, data)
    shutil.move(tmp, out_path)


def embed_audio(deck_path, mp3_path, slide_number, out_path):
    existing = find_slide_audio_target(deck_path, slide_number)
    if existing:
        # The slide already has a wired-up audio object; swap the audio bytes.
        replace_media_part(deck_path, existing, mp3_path, out_path)
        print(f"Replaced existing audio part {existing}")
        return out_path

    prs = Presentation(deck_path)
    slide = prs.slides[slide_number - 1]
    slide.shapes.add_movie(
        mp3_path,
        left=prs.slide_width - Inches(1.6),
        top=prs.slide_height - Inches(1.6),
        width=Inches(1.1),
        height=Inches(1.1),
        mime_type="audio/mpeg",
    )
    tb = slide.shapes.add_textbox(
        prs.slide_width - Inches(2.4),
        prs.slide_height - Inches(0.5),
        Inches(2.2), Inches(0.4),
    )
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Click to play 2-min pitch"
    run.font.size = Pt(10)
    prs.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--script", required=True, help="pitch script .md file")
    ap.add_argument("--deck", required=True, help="input .pptx")
    ap.add_argument("--voice-id", required=True)
    ap.add_argument("--slide", type=int, default=14, help="1-based slide number")
    ap.add_argument("--out", required=True, help="output .pptx")
    ap.add_argument("--api-key", default=os.environ.get("ELEVENLABS_API_KEY"))
    ap.add_argument("--mp3", help="reuse an existing narration mp3 instead of calling the API")
    args = ap.parse_args()

    if args.mp3 and os.path.exists(args.mp3):
        mp3 = args.mp3
    else:
        if not args.api_key:
            sys.exit("Set ELEVENLABS_API_KEY or pass --api-key")
        narration = extract_narration(args.script)
        print(f"Narration: {len(narration)} characters")
        mp3 = args.mp3 or (os.path.splitext(args.out)[0] + ".mp3")
        synthesize(narration, args.voice_id, args.api_key, mp3)
        print(f"Saved narration audio: {mp3} ({os.path.getsize(mp3)} bytes)")

    embed_audio(args.deck, mp3, args.slide, args.out)
    print(f"Embedded audio on slide {args.slide}: {args.out}")


if __name__ == "__main__":
    main()
