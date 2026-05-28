#!/usr/bin/env python3
"""
Re-embed videos with proper poster images (thumbnails) from screenshots
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
import os

PPTX_PATH = '/Users/moderndavinci/Desktop/lightning-pass-beach-club-pitch.pptx'
SCREENSHOTS = '/Users/moderndavinci/Desktop/Applications/Lightning-Project/Lightning-App/screenshots'
VIDEOS = f'{SCREENSHOTS}/videos'

# Colors
NEON_GREEN = RGBColor(57, 255, 20)
GRAY = RGBColor(150, 150, 150)

def remove_videos_from_slide(slide):
    """Remove existing video shapes from slide"""
    shapes_to_remove = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
            shapes_to_remove.append(shape)

    for shape in shapes_to_remove:
        sp = shape._element
        sp.getparent().remove(sp)

def add_video_with_poster(slide, video_path, poster_path, left, top, width, height, label=None):
    """Add video with a proper poster image"""

    # Add the video
    movie = slide.shapes.add_movie(
        video_path,
        left, top, width, height,
        poster_frame_image=poster_path,
        mime_type='video/mp4'
    )

    # Add label if provided
    if label:
        lbl = slide.shapes.add_textbox(left, top + height + Inches(0.05), width, Inches(0.25))
        tf = lbl.text_frame
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(10)
        p.font.color.rgb = NEON_GREEN
        p.alignment = PP_ALIGN.CENTER

    return movie

def main():
    print("\n🎬 Re-embedding videos with proper poster images...\n")

    prs = Presentation(PPTX_PATH)

    # Video configurations: (slide_index, video_file, poster_file, position, size, label)
    video_configs = [
        {
            'slide': 15,  # Slide 16
            'video': 'customer-flow.mp4',
            'poster': 'flow-01-landing.png',
            'pos': (Inches(9.2), Inches(1.4)),
            'size': (Inches(3.5), Inches(5)),
            'label': '▶ Click to Play'
        },
        {
            'slide': 16,  # Slide 17
            'video': 'dj-experience.mp4',
            'poster': 'dj-dashboard.png',
            'pos': (Inches(9.5), Inches(3)),
            'size': (Inches(3.3), Inches(1.9)),
            'label': '▶ DJ Demo'
        },
        {
            'slide': 17,  # Slide 18
            'video': 'venue-dashboard.mp4',
            'poster': 'venue-01-overview.png',
            'pos': (Inches(9.5), Inches(5.5)),
            'size': (Inches(3.3), Inches(1.2)),
            'label': '▶ Dashboard'
        },
        {
            'slide': 18,  # Slide 19
            'video': 'white-label-demo.mp4',
            'poster': 'brand-lightning.png',
            'pos': (Inches(8.8), Inches(5.2)),
            'size': (Inches(4), Inches(1.5)),
            'label': '▶ Branding Demo'
        },
        {
            'slide': 19,  # Slide 20
            'video': 'bouncer-scan.mp4',
            'poster': 'venue-08-scanner.png',
            'pos': (Inches(10.5), Inches(1.4)),
            'size': (Inches(2.3), Inches(1.3)),
            'label': '▶ Scan Demo'
        },
    ]

    for config in video_configs:
        slide = prs.slides[config['slide']]
        video_path = os.path.join(VIDEOS, config['video'])
        poster_path = os.path.join(SCREENSHOTS, config['poster'])

        if not os.path.exists(video_path):
            print(f"  ⚠️  {config['video']} not found")
            continue

        if not os.path.exists(poster_path):
            print(f"  ⚠️  {config['poster']} not found")
            continue

        # Remove existing video
        remove_videos_from_slide(slide)

        # Add video with poster
        add_video_with_poster(
            slide,
            video_path,
            poster_path,
            config['pos'][0], config['pos'][1],
            config['size'][0], config['size'][1],
            config['label']
        )

        print(f"  ✅ Slide {config['slide'] + 1}: {config['video']} with {config['poster']} poster")

    prs.save(PPTX_PATH)
    print(f"\n💾 Saved to {PPTX_PATH}")

    size = os.path.getsize(PPTX_PATH) / 1024 / 1024
    print(f"   File size: {size:.1f} MB\n")

if __name__ == "__main__":
    main()
